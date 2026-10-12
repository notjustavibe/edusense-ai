"""EduSense AI - face recognition attendance server.

Run:  python app.py            (listens on 0.0.0.0:5000)
Env:  EDUSENSE_KEY        optional API key (clients send X-API-Key header)
      EDUSENSE_TOLERANCE  match threshold, lower = stricter (default 0.5)
      EDUSENSE_DB         path to the SQLite file (default: ./edusense.db)
"""
import csv
import io
import json
import os
import re
import sqlite3
import zipfile
from datetime import datetime

import face_recognition
import numpy as np
from flask import Flask, Response, g, jsonify, request, send_from_directory
from PIL import Image, ImageOps

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("EDUSENSE_DB", os.path.join(BASE, "edusense.db"))
TOLERANCE = float(os.environ.get("EDUSENSE_TOLERANCE", "0.5"))
API_KEY = os.environ.get("EDUSENSE_KEY", "")
MAX_SIDE = 1600  # phone photos are downscaled to this before detection

SCHEMA = """
CREATE TABLE IF NOT EXISTS students (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    roll_no TEXT NOT NULL UNIQUE,
    name TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS encodings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    vector TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS sessions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS attendance (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    session_id INTEGER NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    marked_at TEXT NOT NULL,
    distance REAL,
    method TEXT NOT NULL DEFAULT 'face',
    UNIQUE(session_id, student_id)
);
CREATE TABLE IF NOT EXISTS subjects (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    code TEXT NOT NULL UNIQUE,
    name TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS assessments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
    kind TEXT NOT NULL CHECK (kind IN ('quiz', 'assignment', 'mid', 'final')),
    title TEXT NOT NULL,
    max_marks REAL NOT NULL CHECK (max_marks > 0)
);
CREATE TABLE IF NOT EXISTS marks (
    assessment_id INTEGER NOT NULL REFERENCES assessments(id) ON DELETE CASCADE,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    marks REAL NOT NULL,
    PRIMARY KEY (assessment_id, student_id)
);
"""

KIND_ORDER = {"quiz": 0, "assignment": 1, "mid": 2, "final": 3}
KIND_LABEL = {"quiz": "Quiz", "assignment": "Assignment", "mid": "Mid-term", "final": "Final"}

app = Flask(__name__)


# ----------------------------------------------------------------- database
def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


def init_db():
    """Create tables and upgrade an older database in place (no data is lost)."""
    with sqlite3.connect(DB_PATH) as conn:
        conn.executescript(SCHEMA)
        cols = {r[1] for r in conn.execute("PRAGMA table_info(students)")}
        for col, ddl in (
            ("email", "TEXT NOT NULL DEFAULT ''"),
            ("phone", "TEXT NOT NULL DEFAULT ''"),
            ("program", "TEXT NOT NULL DEFAULT ''"),
            ("photo", "BLOB"),
        ):
            if col not in cols:
                conn.execute("ALTER TABLE students ADD COLUMN %s %s" % (col, ddl))
        if not conn.execute("SELECT 1 FROM subjects LIMIT 1").fetchone():
            conn.execute("INSERT INTO subjects (code, name) VALUES ('GEN', 'General')")
        default_subject = conn.execute("SELECT id FROM subjects ORDER BY id LIMIT 1").fetchone()[0]
        scols = {r[1] for r in conn.execute("PRAGMA table_info(sessions)")}
        if "subject_id" not in scols:
            conn.execute("ALTER TABLE sessions ADD COLUMN subject_id INTEGER REFERENCES subjects(id)")
        conn.execute("UPDATE sessions SET subject_id = ? WHERE subject_id IS NULL", (default_subject,))


def now():
    return datetime.now().isoformat(timespec="seconds")


# --------------------------------------------------------------------- auth
@app.before_request
def check_key():
    if not API_KEY or not request.path.startswith("/api/") or request.path == "/api/health":
        return None
    supplied = request.headers.get("X-API-Key") or request.args.get("key")
    if supplied != API_KEY:
        return jsonify(error="Invalid or missing API key"), 401
    return None


def fail(message, code=400):
    return jsonify(error=message), code


# ------------------------------------------------------------- face helpers
def load_image(file_storage):
    """Open an uploaded photo as an RGB numpy array, fixing phone rotation."""
    img = Image.open(file_storage.stream)
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((MAX_SIDE, MAX_SIDE))
    return np.array(img)


def face_thumbnail(arr, location, size=220):
    """JPEG bytes of the detected face (with some margin) for the student profile."""
    top, right, bottom, left = location
    h, w = arr.shape[:2]
    pad = int(0.35 * max(bottom - top, right - left))
    crop = arr[max(0, top - pad): min(h, bottom + pad), max(0, left - pad): min(w, right + pad)]
    img = Image.fromarray(crop)
    img.thumbnail((size, size))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf.getvalue()


def detect(arr):
    locations = face_recognition.face_locations(arr, model="hog")
    encodings = face_recognition.face_encodings(arr, locations)
    return locations, encodings


def known_faces():
    rows = db().execute("SELECT student_id, vector FROM encodings").fetchall()
    ids = [r["student_id"] for r in rows]
    vectors = np.array([json.loads(r["vector"]) for r in rows]) if rows else np.empty((0, 128))
    return ids, vectors


def get_session(sid):
    return db().execute(
        """SELECT s.*, sub.code AS subject_code, sub.name AS subject_name
           FROM sessions s JOIN subjects sub ON sub.id = s.subject_id WHERE s.id = ?""",
        (sid,),
    ).fetchone()


# -------------------------------------------------------------------- routes
@app.get("/api/health")
def health():
    n = db().execute("SELECT COUNT(*) FROM students").fetchone()[0]
    return jsonify(ok=True, students=n, tolerance=TOLERANCE, min_attendance=MIN_ATTENDANCE)


def pct(part, whole):
    return round(100.0 * part / whole, 1) if whole else None


@app.get("/api/students")
def list_students():
    conn = db()
    rows = conn.execute(
        """SELECT s.id, s.roll_no, s.name, s.email, s.phone, s.program,
                  (s.photo IS NOT NULL) AS has_photo, COUNT(e.id) AS samples
           FROM students s LEFT JOIN encodings e ON e.student_id = s.id
           GROUP BY s.id ORDER BY s.roll_no"""
    ).fetchall()
    total_sessions = conn.execute("SELECT COUNT(*) FROM sessions").fetchone()[0]
    present = dict(
        conn.execute("SELECT student_id, COUNT(*) FROM attendance GROUP BY student_id").fetchall()
    )
    out = []
    for r in rows:
        d = dict(r)
        d["present"] = present.get(d["id"], 0)
        d["attendance_pct"] = pct(d["present"], total_sessions)
        out.append(d)
    return jsonify(students=out, total_sessions=total_sessions)


@app.post("/api/students")
def enroll_student():
    roll_no = (request.form.get("roll_no") or "").strip()
    name = (request.form.get("name") or "").strip()
    files = request.files.getlist("image")
    if not roll_no or not name:
        return fail("roll_no and name are required")
    if not files:
        return fail("Attach at least one photo (field name: image)")

    vectors = []
    thumb = None
    for i, f in enumerate(files, start=1):
        try:
            arr = load_image(f)
        except Exception:
            return fail(f"Photo {i} is not a readable image")
        locs, encs = detect(arr)
        if len(encs) != 1:
            return fail(f"Photo {i} must contain exactly one face (found {len(encs)})")
        vectors.append(encs[0])
        if thumb is None:
            thumb = face_thumbnail(arr, locs[0])

    conn = db()
    row = conn.execute("SELECT id FROM students WHERE roll_no = ?", (roll_no,)).fetchone()
    extra = {k: (request.form.get(k) or "").strip() for k in ("email", "phone", "program")}
    if row:
        student_id = row["id"]
        conn.execute("UPDATE students SET name = ? WHERE id = ?", (name, student_id))
        for k, v in extra.items():
            if v:
                conn.execute("UPDATE students SET %s = ? WHERE id = ?" % k, (v, student_id))
        conn.execute(
            "UPDATE students SET photo = COALESCE(photo, ?) WHERE id = ?", (thumb, student_id)
        )
    else:
        student_id = conn.execute(
            "INSERT INTO students (roll_no, name, email, phone, program, photo) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (roll_no, name, extra["email"], extra["phone"], extra["program"], thumb),
        ).lastrowid
    conn.executemany(
        "INSERT INTO encodings (student_id, vector) VALUES (?, ?)",
        [(student_id, json.dumps(v.tolist())) for v in vectors],
    )
    conn.commit()
    return jsonify(id=student_id, roll_no=roll_no, name=name, added_samples=len(vectors)), 201


@app.delete("/api/students/<int:student_id>")
def delete_student(student_id):
    conn = db()
    cur = conn.execute("DELETE FROM students WHERE id = ?", (student_id,))
    conn.commit()
    if cur.rowcount == 0:
        return fail("No such student", 404)
    return jsonify(ok=True)


@app.get("/api/sessions")
def list_sessions():
    """Newest first. ?subject_id=N filters; ?limit=N (default 30, max 500)."""
    subject_id = request.args.get("subject_id", type=int)
    limit = min(request.args.get("limit", 30, type=int), 500)
    rows = db().execute(
        """SELECT s.id, s.title, s.created_at, s.subject_id,
                  sub.code AS subject_code, sub.name AS subject_name,
                  COUNT(a.id) AS present
           FROM sessions s
           JOIN subjects sub ON sub.id = s.subject_id
           LEFT JOIN attendance a ON a.session_id = s.id
           WHERE (? IS NULL OR s.subject_id = ?)
           GROUP BY s.id ORDER BY s.id DESC LIMIT ?""",
        (subject_id, subject_id, limit),
    ).fetchall()
    return jsonify(sessions=[dict(r) for r in rows])


@app.post("/api/sessions")
def create_session():
    body = request.get_json(silent=True) or {}
    title = (body.get("title") or "").strip()
    if not title:
        return fail("title is required")
    conn = db()
    subject_id = body.get("subject_id")
    if subject_id is None:
        subject_id = conn.execute("SELECT id FROM subjects ORDER BY id LIMIT 1").fetchone()[0]
    if not conn.execute("SELECT 1 FROM subjects WHERE id = ?", (subject_id,)).fetchone():
        return fail("No such subject", 404)
    sid = conn.execute(
        "INSERT INTO sessions (title, created_at, subject_id) VALUES (?, ?, ?)",
        (title, now(), subject_id),
    ).lastrowid
    conn.commit()
    return jsonify(id=sid, title=title, subject_id=subject_id), 201


def attendance_payload(sid):
    conn = db()
    present = conn.execute(
        """SELECT s.id AS student_id, s.roll_no, s.name, a.marked_at, a.method, a.distance
           FROM attendance a JOIN students s ON s.id = a.student_id
           WHERE a.session_id = ? ORDER BY s.roll_no""",
        (sid,),
    ).fetchall()
    absent = conn.execute(
        """SELECT id AS student_id, roll_no, name FROM students
           WHERE id NOT IN (SELECT student_id FROM attendance WHERE session_id = ?)
           ORDER BY roll_no""",
        (sid,),
    ).fetchall()
    return [dict(r) for r in present], [dict(r) for r in absent]


@app.get("/api/sessions/<int:sid>/attendance")
def session_attendance(sid):
    session = get_session(sid)
    if not session:
        return fail("No such session", 404)
    present, absent = attendance_payload(sid)
    return jsonify(session=dict(session), present=present, absent=absent)


@app.post("/api/sessions/<int:sid>/recognize")
def recognize(sid):
    if not get_session(sid):
        return fail("No such session", 404)
    f = request.files.get("image")
    if f is None:
        return fail("Attach a photo (field name: image)")
    try:
        arr = load_image(f)
    except Exception:
        return fail("Photo is not a readable image")

    _locs, encs = detect(arr)
    ids, vectors = known_faces()

    best = {}  # student_id -> smallest distance seen in this photo
    unknown = 0
    for enc in encs:
        if len(ids) == 0:
            unknown += 1
            continue
        distances = face_recognition.face_distance(vectors, enc)
        i = int(np.argmin(distances))
        if distances[i] <= TOLERANCE:
            student_id = ids[i]
            best[student_id] = min(best.get(student_id, 9.0), float(distances[i]))
        else:
            unknown += 1

    conn = db()
    recognized = []
    for student_id, dist in best.items():
        cur = conn.execute(
            """INSERT OR IGNORE INTO attendance
               (session_id, student_id, marked_at, distance, method)
               VALUES (?, ?, ?, ?, 'face')""",
            (sid, student_id, now(), dist),
        )
        s = conn.execute(
            "SELECT roll_no, name FROM students WHERE id = ?", (student_id,)
        ).fetchone()
        recognized.append(
            {
                "student_id": student_id,
                "roll_no": s["roll_no"],
                "name": s["name"],
                "distance": round(dist, 3),
                "new": cur.rowcount == 1,
            }
        )
    conn.commit()
    return jsonify(faces_detected=len(encs), recognized=recognized, unknown=unknown)


@app.post("/api/sessions/<int:sid>/mark")
def manual_mark(sid):
    """Manual override: {"student_id": 3, "present": true|false}."""
    if not get_session(sid):
        return fail("No such session", 404)
    body = request.get_json(silent=True) or {}
    student_id = body.get("student_id")
    if not isinstance(student_id, int):
        return fail("student_id (integer) is required")
    conn = db()
    if body.get("present", True):
        conn.execute(
            """INSERT OR IGNORE INTO attendance
               (session_id, student_id, marked_at, method) VALUES (?, ?, ?, 'manual')""",
            (sid, student_id, now()),
        )
    else:
        conn.execute(
            "DELETE FROM attendance WHERE session_id = ? AND student_id = ?", (sid, student_id)
        )
    conn.commit()
    return jsonify(ok=True)


@app.get("/api/sessions/<int:sid>/export.csv")
def export_csv(sid):
    session = get_session(sid)
    if not session:
        return fail("No such session", 404)
    present, absent = attendance_payload(sid)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["roll_no", "name", "status", "marked_at", "method"])
    for r in present:
        writer.writerow([r["roll_no"], r["name"], "present", r["marked_at"], r["method"]])
    for r in absent:
        writer.writerow([r["roll_no"], r["name"], "absent", "", ""])
    return Response(
        out.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="attendance_{sid}.csv"'},
    )


# ============================================================ subjects & marks
def subject_rows():
    return db().execute("SELECT * FROM subjects ORDER BY code").fetchall()


def sorted_assessments(rows):
    return sorted((dict(r) for r in rows), key=lambda r: (KIND_ORDER[r["kind"]], r["id"]))


@app.get("/api/subjects")
def list_subjects():
    rows = db().execute(
        """SELECT sub.id, sub.code, sub.name,
                  (SELECT COUNT(*) FROM sessions WHERE subject_id = sub.id) AS sessions,
                  (SELECT COUNT(*) FROM assessments WHERE subject_id = sub.id) AS assessments
           FROM subjects sub ORDER BY sub.code"""
    ).fetchall()
    return jsonify(subjects=[dict(r) for r in rows])


@app.post("/api/subjects")
def create_subject():
    body = request.get_json(silent=True) or {}
    code = (body.get("code") or "").strip().upper()
    name = (body.get("name") or "").strip()
    if not code or not name:
        return fail("code and name are required")
    conn = db()
    try:
        sid = conn.execute("INSERT INTO subjects (code, name) VALUES (?, ?)", (code, name)).lastrowid
    except sqlite3.IntegrityError:
        return fail("A subject with code %s already exists" % code, 409)
    conn.commit()
    return jsonify(id=sid, code=code, name=name), 201


@app.delete("/api/subjects/<int:subject_id>")
def delete_subject(subject_id):
    conn = db()
    if not conn.execute("SELECT 1 FROM subjects WHERE id = ?", (subject_id,)).fetchone():
        return fail("No such subject", 404)
    used = conn.execute(
        "SELECT (SELECT COUNT(*) FROM sessions WHERE subject_id = ?)"
        " + (SELECT COUNT(*) FROM assessments WHERE subject_id = ?)",
        (subject_id, subject_id),
    ).fetchone()[0]
    if used:
        return fail("This subject still has sessions or assessments; delete those first", 409)
    conn.execute("DELETE FROM subjects WHERE id = ?", (subject_id,))
    conn.commit()
    return jsonify(ok=True)


@app.get("/api/subjects/<int:subject_id>/assessments")
def list_assessments(subject_id):
    rows = db().execute(
        "SELECT id, subject_id, kind, title, max_marks FROM assessments WHERE subject_id = ?",
        (subject_id,),
    ).fetchall()
    return jsonify(assessments=sorted_assessments(rows))


@app.post("/api/subjects/<int:subject_id>/assessments")
def create_assessment(subject_id):
    body = request.get_json(silent=True) or {}
    kind = (body.get("kind") or "").strip().lower()
    title = (body.get("title") or "").strip()
    try:
        max_marks = float(body.get("max_marks"))
    except (TypeError, ValueError):
        return fail("max_marks must be a number")
    if kind not in KIND_ORDER:
        return fail("kind must be one of: quiz, assignment, mid, final")
    if max_marks <= 0:
        return fail("max_marks must be greater than 0")
    conn = db()
    if not conn.execute("SELECT 1 FROM subjects WHERE id = ?", (subject_id,)).fetchone():
        return fail("No such subject", 404)
    if not title:
        n = conn.execute(
            "SELECT COUNT(*) FROM assessments WHERE subject_id = ? AND kind = ?", (subject_id, kind)
        ).fetchone()[0]
        title = KIND_LABEL[kind] + ("" if kind in ("mid", "final") and n == 0 else " %d" % (n + 1))
    aid = conn.execute(
        "INSERT INTO assessments (subject_id, kind, title, max_marks) VALUES (?, ?, ?, ?)",
        (subject_id, kind, title, max_marks),
    ).lastrowid
    conn.commit()
    return jsonify(id=aid, subject_id=subject_id, kind=kind, title=title, max_marks=max_marks), 201


@app.delete("/api/assessments/<int:assessment_id>")
def delete_assessment(assessment_id):
    conn = db()
    cur = conn.execute("DELETE FROM assessments WHERE id = ?", (assessment_id,))
    conn.commit()
    if cur.rowcount == 0:
        return fail("No such assessment", 404)
    return jsonify(ok=True)


@app.get("/api/assessments/<int:assessment_id>/marks")
def get_marks(assessment_id):
    conn = db()
    a = conn.execute("SELECT * FROM assessments WHERE id = ?", (assessment_id,)).fetchone()
    if not a:
        return fail("No such assessment", 404)
    rows = conn.execute(
        """SELECT s.id AS student_id, s.roll_no, s.name, m.marks
           FROM students s LEFT JOIN marks m
             ON m.student_id = s.id AND m.assessment_id = ?
           ORDER BY s.roll_no""",
        (assessment_id,),
    ).fetchall()
    return jsonify(assessment=dict(a), marks=[dict(r) for r in rows])


@app.put("/api/assessments/<int:assessment_id>/marks")
def set_marks(assessment_id):
    """Body: {"marks": {"<student_id>": 8.5, "<student_id>": null}}  (null = clear)."""
    conn = db()
    a = conn.execute("SELECT * FROM assessments WHERE id = ?", (assessment_id,)).fetchone()
    if not a:
        return fail("No such assessment", 404)
    incoming = (request.get_json(silent=True) or {}).get("marks")
    if not isinstance(incoming, dict):
        return fail('Send {"marks": {"<student_id>": value, ...}}')
    parsed = {}
    for key, value in incoming.items():
        try:
            student_id = int(key)
        except ValueError:
            return fail("Bad student id: %r" % key)
        if value is None or (isinstance(value, str) and not value.strip()):
            parsed[student_id] = None
            continue
        try:
            number = float(value)
        except (TypeError, ValueError):
            return fail("Marks for student %d is not a number" % student_id)
        if not 0 <= number <= a["max_marks"]:
            return fail("Marks for student %d must be between 0 and %g" % (student_id, a["max_marks"]))
        parsed[student_id] = number
    known = {r[0] for r in conn.execute("SELECT id FROM students")}
    unknown = [i for i in parsed if i not in known]
    if unknown:
        return fail("Unknown student id(s): %s" % ", ".join(map(str, unknown)), 404)
    for student_id, number in parsed.items():
        if number is None:
            conn.execute(
                "DELETE FROM marks WHERE assessment_id = ? AND student_id = ?",
                (assessment_id, student_id),
            )
        else:
            conn.execute(
                "INSERT INTO marks (assessment_id, student_id, marks) VALUES (?, ?, ?) "
                "ON CONFLICT(assessment_id, student_id) DO UPDATE SET marks = excluded.marks",
                (assessment_id, student_id, number),
            )
    conn.commit()
    return jsonify(ok=True, saved=len(parsed))


# ================================================================== records
MIN_ATTENDANCE = float(os.environ.get("EDUSENSE_MIN_ATTENDANCE", "75"))


def student_record(student_id, subjects=None):
    """Profile + per-subject attendance and marks for one student (None if unknown)."""
    conn = db()
    s = conn.execute(
        "SELECT id, roll_no, name, email, phone, program, (photo IS NOT NULL) AS has_photo "
        "FROM students WHERE id = ?", (student_id,)
    ).fetchone()
    if not s:
        return None
    rec = dict(s)
    rec["subjects"] = []
    all_sessions = all_present = 0
    for sub in subjects if subjects is not None else subject_rows():
        sessions = conn.execute(
            "SELECT COUNT(*) FROM sessions WHERE subject_id = ?", (sub["id"],)
        ).fetchone()[0]
        present = conn.execute(
            """SELECT COUNT(*) FROM attendance a JOIN sessions se ON se.id = a.session_id
               WHERE se.subject_id = ? AND a.student_id = ?""",
            (sub["id"], student_id),
        ).fetchone()[0]
        rows = conn.execute(
            """SELECT a.id, a.kind, a.title, a.max_marks, m.marks
               FROM assessments a
               LEFT JOIN marks m ON m.assessment_id = a.id AND m.student_id = ?
               WHERE a.subject_id = ?""",
            (student_id, sub["id"]),
        ).fetchall()
        assessments = sorted_assessments(rows)
        graded = [a for a in assessments if a["marks"] is not None]
        obtained = sum(a["marks"] for a in graded)
        out_of = sum(a["max_marks"] for a in graded)
        all_sessions += sessions
        all_present += present
        rec["subjects"].append(
            {
                "id": sub["id"],
                "code": sub["code"],
                "name": sub["name"],
                "sessions": sessions,
                "present": present,
                "attendance_pct": pct(present, sessions),
                "assessments": assessments,
                "obtained": obtained,
                "out_of": out_of,
                "marks_pct": pct(obtained, out_of),
            }
        )
    rec["sessions"] = all_sessions
    rec["present"] = all_present
    rec["attendance_pct"] = pct(all_present, all_sessions)
    return rec


@app.get("/api/students/<int:student_id>")
def get_student(student_id):
    rec = student_record(student_id)
    if rec is None:
        return fail("No such student", 404)
    rec["min_attendance"] = MIN_ATTENDANCE
    return jsonify(rec)


@app.put("/api/students/<int:student_id>")
def update_student(student_id):
    body = request.get_json(silent=True) or {}
    conn = db()
    if not conn.execute("SELECT 1 FROM students WHERE id = ?", (student_id,)).fetchone():
        return fail("No such student", 404)
    for field in ("roll_no", "name", "email", "phone", "program"):
        if field in body:
            value = (body[field] or "").strip()
            if field in ("roll_no", "name") and not value:
                return fail("%s cannot be empty" % field)
            try:
                conn.execute("UPDATE students SET %s = ? WHERE id = ?" % field, (value, student_id))
            except sqlite3.IntegrityError:
                conn.rollback()
                return fail("Another student already has roll number %s" % value, 409)
    conn.commit()
    return jsonify(ok=True)


@app.get("/api/students/<int:student_id>/photo")
def student_photo(student_id):
    row = db().execute("SELECT photo FROM students WHERE id = ?", (student_id,)).fetchone()
    if not row or not row["photo"]:
        return fail("No photo", 404)
    return Response(row["photo"], mimetype="image/jpeg", headers={"Cache-Control": "max-age=60"})


# ================================================================== exports
def fmt_num(x):
    return ("%g" % x) if x is not None else ""


def safe_name(text):
    return re.sub(r"[^\w\-]+", "_", text).strip("_") or "student"


def download(data, mimetype, filename):
    return Response(
        data,
        mimetype=mimetype,
        headers={"Content-Disposition": 'attachment; filename="%s"' % filename},
    )


def build_pdf(rec, photo):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
        bottomMargin=16 * mm, title="Student report - %s" % rec["name"], author="EduSense AI",
    )
    story = [
        Paragraph("Student Report", styles["Title"]),
        Paragraph("EduSense AI &middot; generated %s" % datetime.now().strftime("%d %b %Y"),
                  styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    profile = [
        ["Roll number", rec["roll_no"]], ["Name", rec["name"]],
        ["Program", rec["program"] or "-"], ["Email", rec["email"] or "-"],
        ["Phone", rec["phone"] or "-"],
    ]
    ptable = Table(profile, colWidths=[30 * mm, 80 * mm], hAlign="LEFT")
    ptable.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#555555")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    if photo:
        img = RLImage(io.BytesIO(photo), width=32 * mm, height=32 * mm, kind="proportional")
        head = Table([[img, ptable]], colWidths=[38 * mm, 115 * mm])
        head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(head)
    else:
        story.append(ptable)
    story.append(Spacer(1, 5 * mm))

    overall = rec["attendance_pct"]
    low = overall is not None and overall < MIN_ATTENDANCE
    story.append(Paragraph(
        "<b>Overall attendance:</b> %s of %s sessions (%s)" % (
            rec["present"], rec["sessions"], "%.1f%%" % overall if overall is not None else "n/a")
        + (' &nbsp;<font color="#c0392b"><b>below %g%%</b></font>' % MIN_ATTENDANCE if low else ""),
        styles["Normal"]))

    for sub in rec["subjects"]:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("%s &ndash; %s" % (sub["code"], sub["name"]), styles["Heading2"]))
        a = sub["attendance_pct"]
        flag = a is not None and a < MIN_ATTENDANCE
        story.append(Paragraph(
            "Attendance: %s of %s sessions (%s)" % (
                sub["present"], sub["sessions"], "%.1f%%" % a if a is not None else "n/a")
            + (' &nbsp;<font color="#c0392b"><b>below %g%%</b></font>' % MIN_ATTENDANCE if flag else ""),
            styles["Normal"]))
        story.append(Spacer(1, 2 * mm))
        rows = [["Assessment", "Type", "Marks", "Out of", "%"]]
        for x in sub["assessments"]:
            got = x["marks"]
            rows.append([x["title"], KIND_LABEL[x["kind"]], fmt_num(got) if got is not None else "-",
                         fmt_num(x["max_marks"]),
                         "%.1f" % (100 * got / x["max_marks"]) if got is not None else "-"])
        if len(rows) == 1:
            story.append(Paragraph("<i>No assessments set up for this subject yet.</i>",
                                   styles["Normal"]))
            continue
        rows.append(["Total (graded)", "", fmt_num(sub["obtained"]), fmt_num(sub["out_of"]),
                     "%.1f" % sub["marks_pct"] if sub["marks_pct"] is not None else "-"])
        t = Table(rows, colWidths=[58 * mm, 32 * mm, 22 * mm, 22 * mm, 20 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#eef2f7")),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c8d0da")),
        ]))
        story.append(t)
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph(
        "<font size=8 color='#777777'>Marks percentage counts graded assessments only. "
        "Attendance is the share of recorded class sessions in which the student was marked "
        "present.</font>", styles["Normal"]))
    doc.build(story)
    return buf.getvalue()


def student_pdf(student_id):
    rec = student_record(student_id)
    if rec is None:
        return None, None
    photo = db().execute("SELECT photo FROM students WHERE id = ?", (student_id,)).fetchone()[0]
    return rec, build_pdf(rec, photo)


@app.get("/api/students/<int:student_id>/report.pdf")
def report_pdf(student_id):
    rec, pdf = student_pdf(student_id)
    if rec is None:
        return fail("No such student", 404)
    return download(pdf, "application/pdf", "%s_%s.pdf" % (safe_name(rec["roll_no"]), safe_name(rec["name"])))


@app.get("/api/export/reports.zip")
def reports_zip():
    ids = [r[0] for r in db().execute("SELECT id FROM students ORDER BY roll_no")]
    if not ids:
        return fail("No students enrolled yet", 404)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for sid in ids:
            rec, pdf = student_pdf(sid)
            z.writestr("%s_%s.pdf" % (safe_name(rec["roll_no"]), safe_name(rec["name"])), pdf)
    return download(buf.getvalue(), "application/zip", "edusense_student_reports.zip")


@app.get("/api/export/all.xlsx")
def export_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    subjects = subject_rows()
    ids = [r[0] for r in db().execute("SELECT id FROM students ORDER BY roll_no")]
    records = [student_record(i, subjects) for i in ids]

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    red = Font(color="C0392B", bold=True)

    def finish(ws, header_row=1):
        for cell in ws[header_row]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
        for col in range(1, ws.max_column + 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in ws[get_column_letter(col)])
            ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 40)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    header = ["Roll No", "Name", "Program", "Email", "Phone", "Overall Attendance %"]
    for sub in subjects:
        header += ["%s Attendance %%" % sub["code"], "%s Marks %%" % sub["code"]]
    ws.append(header)
    for r in records:
        row = [r["roll_no"], r["name"], r["program"], r["email"], r["phone"], r["attendance_pct"]]
        for s in r["subjects"]:
            row += [s["attendance_pct"], s["marks_pct"]]
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            h = header[cell.column - 1]
            if "Attendance" in h and isinstance(cell.value, (int, float)) and cell.value < MIN_ATTENDANCE:
                cell.font = red
    finish(ws)

    used = {"Summary"}
    for sub in subjects:
        title = re.sub(r"[\[\]:*?/\\]", "", sub["code"])[:31] or "Subject"
        while title in used:
            title = title[:28] + "_%d" % len(used)
        used.add(title)
        sheet = wb.create_sheet(title)
        assessments = sorted_assessments(db().execute(
            "SELECT id, kind, title, max_marks FROM assessments WHERE subject_id = ?", (sub["id"],)))
        sheet.append(["Roll No", "Name", "Present", "Sessions", "Attendance %"]
                     + ["%s (/%s)" % (a["title"], fmt_num(a["max_marks"])) for a in assessments]
                     + ["Obtained", "Out of", "Marks %"])
        for r in records:
            s = next(x for x in r["subjects"] if x["id"] == sub["id"])
            by_id = {a["id"]: a["marks"] for a in s["assessments"]}
            sheet.append([r["roll_no"], r["name"], s["present"], s["sessions"], s["attendance_pct"]]
                         + [by_id.get(a["id"]) for a in assessments]
                         + [s["obtained"] if s["out_of"] else None, s["out_of"] or None, s["marks_pct"]])
            att = sheet.cell(row=sheet.max_row, column=5)
            if isinstance(att.value, (int, float)) and att.value < MIN_ATTENDANCE:
                att.font = red
        finish(sheet)

    buf = io.BytesIO()
    wb.save(buf)
    return download(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "edusense_records_%s.xlsx" % datetime.now().strftime("%Y%m%d"),
    )


# ================================================================= dashboard
@app.get("/")
def dashboard():
    return send_from_directory(os.path.join(BASE, "static"), "dashboard.html")


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
